from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Iterable, Iterator, Optional

from lxml import etree

from .search_url import normalize_name
from .structured_price_lists import json_rows, xml_rows

PARSED_FORMATS = ("xlsx", "csv", "json", "xml")
MAX_HEADER_SEARCH_ROWS = 15
MIN_UNIT_PRICE = 20_000
MIN_PRICE_PER_M2 = 1_000
MAX_PRICE_PER_M2 = 100_000

_EMPTY_VALUES = {"", "x", "-", "--", "brak", "nie dotyczy", "n/d", "nd", "nd.", "none", "null", "0", "0,00", "0.00"}
_PLACE_PREFIX = re.compile(r"^(?:m\.|m\.st\.|miasto|gmina|gm\.|powiat|pow\.|woj\.|wojewodztwo)\s+")
_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PROJECT = ("przedsiewziecia", "inwestycji", "zadania")
_NOT_UNIT_PRICE = ("data", "m2", "uwzgledniajac", "czesci", "pomieszczen", "praw", "swiadczen")
_NOT_UNIT_LABEL = ("adres", "data", "cena", "rodzaj", "siedziby", "oznaczenie")


@dataclass
class PriceListUnit:
    developer: Optional[str]
    investment: Optional[str]
    voivodeship: Optional[str]
    gmina: Optional[str]
    town: str
    street: Optional[str]
    building_number: Optional[str]
    unit: str
    kind: Optional[str]
    area: Optional[float]
    rooms: Optional[int]
    floor: Optional[str]
    price: Optional[int]
    price_per_m2: Optional[int]
    price_since: Optional[date]
    price_with_extras: Optional[int]
    extras: list[str] = field(default_factory=list)
    extras_value: int = 0
    prospectus_url: Optional[str] = None

    @property
    def key(self) -> str:
        return "|".join(place_name(part) or "" for part in (self.town, self.street, self.building_number, self.unit))


@dataclass(frozen=True)
class _Columns:
    developer: Optional[int]
    investment: Optional[int]
    voivodeship: Optional[int]
    gmina: Optional[int]
    town: Optional[int]
    street: Optional[int]
    building_number: Optional[int]
    unit: Optional[int]
    kind: Optional[int]
    area: Optional[int]
    rooms: Optional[int]
    floor: Optional[int]
    price: Optional[int]
    price_per_m2: Optional[int]
    price_since: Optional[int]
    price_with_extras: Optional[int]
    extras: tuple[tuple[Optional[int], Optional[int], Optional[int]], ...]
    prospectus: Optional[int]


def place_name(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    plain = re.sub(r"\s+", " ", normalize_name(str(value)).strip().lower())
    plain = _PLACE_PREFIX.sub("", plain)
    return plain or None


def parse_price_list(content: bytes, file_format: str) -> list[PriceListUnit]:
    file_format = (file_format or "").lower()
    if file_format == "xlsx":
        rows = _xlsx_rows(content)
    elif file_format == "csv":
        rows = _csv_rows(content)
    elif file_format == "json":
        rows = json_rows(content)
    elif file_format == "xml":
        rows = xml_rows(content)
    else:
        return []
    return list(units_from_rows(rows))


def units_from_rows(rows: Iterable[list]) -> Iterator[PriceListUnit]:
    rows = iter(rows)
    columns = None
    for _, row in zip(range(MAX_HEADER_SEARCH_ROWS), rows):
        columns = _columns_of(row)
        if columns is not None:
            break
    if columns is None:
        return
    units: dict[str, PriceListUnit] = {}
    for row in rows:
        unit = _unit_of(row, columns)
        if unit is None:
            continue
        existing = units.get(unit.key)
        if existing is None:
            units[unit.key] = unit
        else:
            for extra in unit.extras:
                if extra not in existing.extras:
                    existing.extras.append(extra)
            existing.extras_value = max(existing.extras_value, unit.extras_value)
    yield from units.values()


def _columns_of(row: list) -> Optional[_Columns]:
    headers = [_header(cell) for cell in row]

    def find(*required: str, excluded: tuple[str, ...] = (), starts: Optional[str] = None) -> Optional[int]:
        for index, header in enumerate(headers):
            if starts is not None and not header.startswith(starts) and f" {starts} " not in f" {header} ":
                continue
            if all(part in header for part in required) and not any(part in header for part in excluded):
                return index
        return None

    def of_project(*required: str, excluded: tuple[str, ...] = ()) -> Optional[int]:
        for project in _PROJECT:
            index = find(*required, project, excluded=("siedziby", "sprzedaz") + excluded)
            if index is not None:
                return index
        return None

    price_per_m2 = find("cena", "m2", excluded=("data",), starts="cena")
    price = _first(
        find("cena", "ilocz", excluded=("data", "uwzgledniajac"), starts="cena"),
        find("cena", excluded=_NOT_UNIT_PRICE, starts="cena"),
    )
    unit = _first(
        find("lokal", "dewelopera", excluded=_NOT_UNIT_LABEL),
        find("nadany przez dewelopera", excluded=_NOT_UNIT_LABEL),
    )
    if price_per_m2 is None or price is None or unit is None:
        return _loose_columns_of(headers)
    return _Columns(
        developer=find("nazwa dewelopera"),
        investment=_first(find("nazwa", "inwestycji", excluded=("dewelopera", "adres strony")), find("nazwa przedsiewziecia")),
        voivodeship=of_project("wojewodztwo"),
        gmina=of_project("gmina"),
        town=of_project("miejscowosc"),
        street=of_project("ulica"),
        building_number=_first(of_project("nr nieruchomosci"), of_project("numer nieruchomosci")),
        unit=unit,
        kind=find("rodzaj nieruchomosci"),
        area=find("powierzchn", excluded=("cena", "data", "przynalezn", "czesci")),
        rooms=_first(find("liczba pokoi"), find("pokoi", excluded=("cena",))),
        floor=_first(find("pietro"), find("kondygnacj", excluded=("cena",))),
        price=price,
        price_per_m2=price_per_m2,
        price_since=find("m2", "data", starts="data"),
        price_with_extras=find("cena", "uwzgledniajac", excluded=("data",), starts="cena"),
        extras=(
            (find("rodzaj czesci"), find("oznaczenie czesci"), find("cena czesci", excluded=("data",))),
            (find("rodzaj pomieszczen"), find("oznaczenie pomieszczen"), find("wyszczegolnienie cen pomieszczen")),
            (find("wyszczegolnienie praw"), None, find("wartosc praw")),
            (find("wyszczegolnienie rodzajow innych"), None, find("wartosc innych")),
        ),
        prospectus=find("prospekt"),
    )


def _loose_columns_of(headers: list[str]) -> Optional[_Columns]:
    def find(*candidates: str, excluded: tuple[str, ...] = ()) -> Optional[int]:
        for candidate in candidates:
            for index, header in enumerate(headers):
                words = f" {header} "
                tail = " ".join(header.split()[-(len(candidate.split()) + 1):])
                if f" {candidate} " in words and header.endswith(candidate.split()[-1])                         and not any(part in tail for part in excluded):
                    return index
        return None

    elsewhere = ("sales", "office", "siedzib", "biur", "sprzedaz", "developer", "dewelopera", "company", "firma",
                 "kontakt", "contact", "korespond", "phone", "telefon", "tel", "fax", "tax", "nip", "regon", "krs",
                 "postal", "pocztow", "email", "mail", "url", "www")
    town = find("miejscowosc", "miasto", "city", "town", "locality", "miejscowosci", excluded=elsewhere)
    unit = find("nr lokalu", "numer lokalu", "id lokalu", "unit number", "unit id", "unit no", "lokal no", "nr mieszkania", "numer mieszkania",
                "oznaczenie lokalu", "lokal", "unit", "product name", "symbol", "nr obiektu", "identyfikator", "numer",
                "number", "nr", "id",
                excluded=elsewhere + ("inwestycj", "investment", "project", "projekt", "budyn", "building", "ulic",
                                      "street", "status", "typ", "type", "rodzaj", "kind", "kondygn", "floor", "pietro",
                                      "pow", "area", "cena", "price"))
    price_words = ("cena", "price", "gross", "brutto", "wartosc", "value")
    not_total = ("m2", "metr", "meter", "data", "date", "promoc", "najnizsz", "lowest", "30 dni", "netto", "net", "vat",
                 "historia", "history", "dzialk", "parking", "garaz", "komork", "miejsc", "postoj")
    price = None
    for preferred in ("cena calkowita", "cena aktualna", "cena brutto", "total price", "price gross", "price brutto",
                      "cena lokalu",
                      "cena mieszkania", "aktualna cena cena", "cena", "price", "gross", "brutto"):
        price = find(preferred, excluded=not_total)
        if price is not None:
            break
    tails = [" ".join(header.split()[-4:]) for header in headers]
    per_m2 = next((index for index, tail in enumerate(tails)
                   if any(word in tail for word in price_words) and any(unit in tail for unit in ("m2", "metr", "meter"))
                   and not any(part in tail for part in ("data", "date", "promoc", "najnizsz", "lowest", "netto",
                                                         "calkowit"))),
                  None)
    if town is None or unit is None or price is None:
        return None
    return _Columns(
        developer=find("developer name", "nazwa dewelopera", "company name", "firma deweloperska", "developer",
                       "deweloper"),
        investment=find("investment name", "nazwa inwestycji", "project name", "nazwa przedsiewziecia",
                        "nazwa projektu", "inwestycja", "investment"),
        voivodeship=find("wojewodztwo", "voivodeship", "region", excluded=elsewhere),
        gmina=find("gmina", "municipality", "commune", excluded=elsewhere),
        town=town,
        street=find("ulica", "street", excluded=elsewhere + ("number", "numer", "nr")),
        building_number=find("numer ulicy", "nr ulicy", "street number", "building number", "buildings number",
                             "numer budynku",
                             "nr budynku", "nr domu", "numer domu", "budynek", "building", excluded=elsewhere),
        unit=unit,
        kind=find("rodzaj nieruchomosci", "typ nieruchomosci", "product type", "rodzaj", "typ", "type", "kind"),
        area=find("powierzchnia uzytkowa", "pow uzytkowa", "powierzchnia m2", "powierzchnia", "metraz m2", "metraz",
                  "usable area", "area m2", "area", "pow u", "powu",
                  excluded=("dzialk", "ogrod", "balkon", "taras", "plot", "garden", "calkowit", "cena", "price",
                            "loggi", "przynalez")),
        rooms=find("liczba pokoi", "ilosc pokoi", "pokoje", "pokoi", "rooms", "rooms count"),
        floor=find("pietro", "kondygnacja", "floor"),
        price=price,
        price_per_m2=per_m2,
        price_since=next((index for index, tail in enumerate(tails)
                          if any(word in tail for word in ("data", "date", "obowiazuje", "valid from"))
                          and not any(part in tail for part in ("histor", "name", "nazwa", "set", "utworz",
                                                                   "created", "generat", "oddania", "odbior"))), None),
        price_with_extras=None,
        extras=(),
        prospectus=find("prospekt", "prospectus", "prospekt informacyjny"),
    )


def _unit_of(row: list, columns: _Columns) -> Optional[PriceListUnit]:
    def cell(index: Optional[int]):
        if index is None or index >= len(row):
            return None
        value = row[index]
        if isinstance(value, str):
            value = value.strip()
            if value.lower() in _EMPTY_VALUES:
                return None
        return value

    town = _text(cell(columns.town))
    unit = _text(cell(columns.unit))
    if not town or not unit:
        return None
    price = _amount(cell(columns.price))
    price_per_m2 = _amount(cell(columns.price_per_m2))
    if price is not None and price < MIN_UNIT_PRICE:
        price = None
    if price_per_m2 is not None and not MIN_PRICE_PER_M2 <= price_per_m2 <= MAX_PRICE_PER_M2:
        price_per_m2 = None
    area = _amount(cell(columns.area), integer=False)
    if area is None and price and price_per_m2:
        area = round(price / price_per_m2, 2)
    extras, extras_value = [], 0
    for kind_index, label_index, price_index in columns.extras:
        kind, label, value = _text(cell(kind_index)), _text(cell(label_index)), _text(cell(price_index))
        if not kind and not value:
            continue
        description = " ".join(part for part in (kind, label) if part)
        extras.append(f"{description}: {value}" if value else description)
        extras_value += sum(_amounts_in(value))
    return PriceListUnit(
        developer=_text(cell(columns.developer)),
        investment=_text(cell(columns.investment)),
        voivodeship=_text(cell(columns.voivodeship)),
        gmina=_text(cell(columns.gmina)),
        town=town,
        street=_text(cell(columns.street)),
        building_number=_text(cell(columns.building_number)),
        unit=unit,
        kind=_text(cell(columns.kind)),
        area=area,
        rooms=_rooms(cell(columns.rooms)),
        floor=_text(cell(columns.floor)),
        price=price,
        price_per_m2=price_per_m2,
        price_since=_date(cell(columns.price_since)),
        price_with_extras=_amount(cell(columns.price_with_extras)),
        extras=extras,
        extras_value=extras_value,
        prospectus_url=_text(cell(columns.prospectus)),
    )


def _first(*indexes: Optional[int]) -> Optional[int]:
    return next((index for index in indexes if index is not None), None)


def _header(cell) -> str:
    plain = normalize_name(str(cell or "")).lower()
    plain = re.sub(r"m\s*[2²]|m\s*\^\s*2|metr\w* kw\w*", "m2", plain)
    return re.sub(r"[^a-z0-9]+", " ", plain).strip()


def _text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value).replace("﻿", "")).strip()
    return text or None


def _rooms(value) -> Optional[int]:
    number = _amount(value)
    return number if number is not None and number <= 20 else None


def _amount(value, integer: bool = True):
    if value is None or isinstance(value, (date, datetime)):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        numbers = _amounts_in(str(value), integer=False)
        if len(numbers) != 1:
            return None
        number = numbers[0]
    if number <= 0:
        return None
    return int(round(number)) if integer else round(number, 2)


def _amounts_in(value: Optional[str], integer: bool = True) -> list:
    if not value:
        return []
    numbers = []
    for match in re.finditer(r"\d[\d\s .,]*", value):
        raw = re.sub(r"[\s ]", "", match.group(0)).rstrip(".,")
        if "," in raw and "." in raw:
            raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
        elif "," in raw:
            whole, _, fraction = raw.rpartition(",")
            raw = f"{whole.replace(',', '')}.{fraction}" if len(fraction) <= 2 else raw.replace(",", "")
        elif raw.count(".") > 1 or (raw.count(".") == 1 and len(raw.rpartition(".")[2]) == 3):
            raw = raw.replace(".", "")
        try:
            number = float(raw)
        except ValueError:
            continue
        numbers.append(int(round(number)) if integer else number)
    return numbers


def _date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    match = re.search(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", text)
    if match:
        year, month, day = (int(part) for part in match.groups())
    else:
        match = re.search(r"(\d{1,2})[-./](\d{1,2})[-./](\d{4})", text)
        if not match:
            return None
        day, month, year = (int(part) for part in match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _xlsx_rows(content: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return _raw_xlsx_rows(content)
    try:
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _raw_xlsx_rows(content: bytes) -> list[list]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        names = archive.namelist()
        shared = []
        if "xl/sharedStrings.xml" in names:
            for item in etree.fromstring(archive.read("xl/sharedStrings.xml")).iterfind(f"{{{_SHEET_NS}}}si"):
                shared.append("".join(item.itertext()))
        sheet = min((name for name in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)), default=None)
        if sheet is None:
            return []
        rows = []
        for row in etree.fromstring(archive.read(sheet)).iter(f"{{{_SHEET_NS}}}row"):
            values: dict[int, object] = {}
            for cell in row.iterfind(f"{{{_SHEET_NS}}}c"):
                column = _column_index(cell.get("r", ""))
                kind = cell.get("t")
                if kind == "inlineStr":
                    values[column] = "".join(cell.itertext())
                    continue
                raw = cell.findtext(f"{{{_SHEET_NS}}}v")
                if raw is None:
                    continue
                values[column] = shared[int(raw)] if kind == "s" else raw
            rows.append([values.get(index) for index in range(max(values, default=-1) + 1)])
        return rows


def _column_index(reference: str) -> int:
    index = 0
    for letter in re.match(r"[A-Z]*", reference).group(0):
        index = index * 26 + ord(letter) - 64
    return index - 1


def _csv_rows(content: bytes) -> Iterator[list]:
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return
    first_line = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t", "|"), key=first_line.count)
    yield from csv.reader(io.StringIO(text), delimiter=delimiter)
